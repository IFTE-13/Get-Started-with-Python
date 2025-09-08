import openpyxl as xl
from openpyxl.chart import BarChart, Reference

def process_workbook(filename, markup_percentage=10):
    # Load workbook and select sheet
    wb = xl.load_workbook(filename)
    sheet = wb['Sheet1']

    # Loop through rows starting from 2 (assuming row 1 is header)
    for row in range(2, sheet.max_row + 1):
        original_price_cell = sheet.cell(row, 3)  # Column C
        original_price = original_price_cell.value

        # Calculate corrected price with markup
        corrected_price = original_price * (1 + markup_percentage / 100)

        # Store corrected price in column D
        corrected_price_cell = sheet.cell(row, 4)  # Column D
        corrected_price_cell.value = round(corrected_price, 2)  # round to 2 decimal places

    # Create a bar chart for corrected prices
    values = Reference(sheet, min_row=2, max_row=sheet.max_row, min_col=4, max_col=4)
    chart = BarChart()
    chart.title = "Corrected Prices"
    chart.y_axis.title = "Price"
    chart.x_axis.title = "Item"
    chart.add_data(values, titles_from_data=False)

    # Add the chart to the sheet
    sheet.add_chart(chart, "E2")

    # Save the workbook
    wb.save(filename)
    print(f"Workbook '{filename}' processed successfully with {markup_percentage}% markup.")

# Example usage:
process_workbook("products.xlsx", markup_percentage=10)
